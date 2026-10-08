from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.tests.apoio_compras import compra_registrada, novo_cliente


def abrir_planilha(resposta):
    return load_workbook(BytesIO(resposta.content))


@pytest.mark.parametrize(
    "recurso",
    ["produtos", "movimentacoes", "fornecedores", "compras", "relatorio-compras"],
)
def test_exportar_em_excel_e_pdf(cliente_autenticado, recurso):
    compra_registrada(cliente_autenticado)

    excel = cliente_autenticado.get(f"/exportacoes/{recurso}?formato=xlsx")
    pdf = cliente_autenticado.get(f"/exportacoes/{recurso}?formato=pdf")

    assert excel.status_code == 200
    assert excel.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f'filename="{recurso}-' in excel.headers["content-disposition"]
    assert excel.headers["content-disposition"].endswith('.xlsx"')
    assert abrir_planilha(excel).sheetnames

    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")


def test_planilha_de_produtos(cliente_autenticado):
    compra_registrada(cliente_autenticado)

    planilha = abrir_planilha(
        cliente_autenticado.get("/exportacoes/produtos?formato=xlsx")
    )
    aba = planilha["Produtos"]
    linhas = list(aba.values)

    assert linhas[0] == (
        "Código",
        "Produto",
        "Categoria",
        "Fornecedor",
        "Quantidade",
        "Preço",
        "Valor em estoque",
    )
    assert {linha[1] for linha in linhas[1:-1]} == {"Cabo", "Conector"}
    assert linhas[-1][1] == "Total"
    assert aba["F2"].number_format == "R$ #,##0.00"


def test_relatorio_de_compras_tem_varias_abas(cliente_autenticado):
    compra_registrada(cliente_autenticado)

    planilha = abrir_planilha(
        cliente_autenticado.get("/exportacoes/relatorio-compras?formato=xlsx")
    )

    assert planilha.sheetnames == ["Resumo", "Por mês", "Por fornecedor", "Por produto"]

    resumo = dict(list(planilha["Resumo"].values)[1:])

    assert resumo["Compras"] == "1"
    assert resumo["Valor comprado"] == "R$ 42,00"


def test_exportar_compras_com_filtros(cliente_autenticado):
    compra_registrada(cliente_autenticado, data_compra="2026-09-15")

    setembro = abrir_planilha(
        cliente_autenticado.get(
            "/exportacoes/compras?formato=xlsx&data_inicio=2026-09-01&data_fim=2026-09-30"
        )
    )
    outubro = abrir_planilha(
        cliente_autenticado.get(
            "/exportacoes/compras?formato=xlsx&data_inicio=2026-10-01&data_fim=2026-10-31"
        )
    )

    assert setembro["Compras"].max_row == 3
    assert outubro["Compras"].max_row == 1


def test_exportacao_vazia_ainda_gera_arquivo(cliente_autenticado):
    excel = cliente_autenticado.get("/exportacoes/compras?formato=xlsx")
    pdf = cliente_autenticado.get("/exportacoes/compras?formato=pdf")

    assert list(abrir_planilha(excel)["Compras"].values)[0][0] == "Data"
    assert pdf.content.startswith(b"%PDF")


def test_formato_e_periodo_invalidos(cliente_autenticado):
    assert cliente_autenticado.get("/exportacoes/produtos?formato=doc").status_code == 422
    assert cliente_autenticado.get(
        "/exportacoes/compras?data_inicio=2026-10-31&data_fim=2026-10-01"
    ).status_code == 400


def test_consulta_pode_exportar(cliente_autenticado):
    resposta = novo_cliente("CONSULTA").get("/exportacoes/produtos?formato=pdf")

    assert resposta.status_code == 200
