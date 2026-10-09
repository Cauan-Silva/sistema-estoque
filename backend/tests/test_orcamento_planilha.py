from io import BytesIO

from openpyxl import load_workbook

from backend.tests.apoio_compras import (
    criar_fornecedor,
    criar_solicitacao_com_dois_itens,
    novo_cliente,
    registrar_cotacao,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def baixar_modelo(cliente, solicitacao_id):
    resposta = cliente.get(f"/solicitacoes-compra/{solicitacao_id}/cotacoes/modelo")

    assert resposta.status_code == 200
    assert resposta.headers["content-type"] == XLSX

    return load_workbook(BytesIO(resposta.content))


def campo(aba, rotulo):
    for linha in aba.iter_rows(min_col=1, max_col=1):
        if linha[0].value == rotulo:
            return aba.cell(row=linha[0].row, column=2)
    raise AssertionError(rotulo)


def linhas_itens(aba):
    for linha in aba.iter_rows(min_col=1, max_col=1):
        if linha[0].value == "Código":
            inicio = linha[0].row + 1
            break
    return {
        aba.cell(row=r, column=1).value: r
        for r in range(inicio, aba.max_row + 1)
        if isinstance(aba.cell(row=r, column=1).value, int)
    }


def preencher(livro, fornecedor, precos, frete=25, prazo=7, validade="31/12/2099"):
    aba = livro["Orçamento"]
    campo(aba, "Fornecedor").value = fornecedor
    campo(aba, "Frete (R$)").value = frete
    campo(aba, "Prazo de entrega (dias)").value = prazo
    campo(aba, "Validade da proposta").value = validade
    linhas = linhas_itens(aba)
    for produto_id, preco in precos.items():
        aba.cell(row=linhas[produto_id], column=4, value=preco)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def importar(cliente, solicitacao_id, conteudo, substituir=False):
    return cliente.post(
        f"/solicitacoes-compra/{solicitacao_id}/cotacoes/importar",
        params={"substituir": substituir},
        files={"arquivo": ("orcamento.xlsx", conteudo, XLSX)},
    )


def test_modelo_traz_itens_da_solicitacao(cliente_autenticado):
    solicitacao_id, cabo, conector = criar_solicitacao_com_dois_itens(cliente_autenticado)

    aba = baixar_modelo(cliente_autenticado, solicitacao_id)["Orçamento"]

    assert campo(aba, "Solicitação Nº").value == solicitacao_id
    assert set(linhas_itens(aba)) == {cabo, conector}


def test_importar_planilha_cria_cotacao(cliente_autenticado):
    solicitacao_id, cabo, conector = criar_solicitacao_com_dois_itens(cliente_autenticado)
    criar_fornecedor(cliente_autenticado, "Atacado São João")

    conteudo = preencher(
        baixar_modelo(cliente_autenticado, solicitacao_id),
        "atacado sao joao",
        {cabo: "12,50", conector: 3},
    )

    resposta = importar(cliente_autenticado, solicitacao_id, conteudo)

    assert resposta.status_code == 201
    dados = resposta.json()
    assert dados["fornecedor"] == "Atacado São João"
    assert dados["frete"] == 25
    assert dados["prazo_entrega_dias"] == 7
    assert dados["validade"] == "2099-12-31"
    assert dados["valor_itens"] == 12.5 * 10 + 3 * 4
    assert dados["valor_total"] == 12.5 * 10 + 3 * 4 + 25


def test_importar_parcial_ignora_itens_sem_preco(cliente_autenticado):
    solicitacao_id, cabo, _ = criar_solicitacao_com_dois_itens(cliente_autenticado)
    criar_fornecedor(cliente_autenticado, "Fornecedor A")

    conteudo = preencher(baixar_modelo(cliente_autenticado, solicitacao_id), "Fornecedor A", {cabo: 10})

    resposta = importar(cliente_autenticado, solicitacao_id, conteudo)

    assert resposta.status_code == 201
    assert [i["produto_id"] for i in resposta.json()["itens"]] == [cabo]


def test_importar_de_novo_pede_para_substituir(cliente_autenticado):
    solicitacao_id, cabo, conector = criar_solicitacao_com_dois_itens(cliente_autenticado)
    criar_fornecedor(cliente_autenticado, "Fornecedor A")
    livro = baixar_modelo(cliente_autenticado, solicitacao_id)

    primeira = importar(cliente_autenticado, solicitacao_id, preencher(livro, "Fornecedor A", {cabo: 10, conector: 2}))
    assert primeira.status_code == 201

    nova = preencher(livro, "Fornecedor A", {cabo: 9, conector: 2}, frete=0)

    duplicada = importar(cliente_autenticado, solicitacao_id, nova)
    assert duplicada.status_code == 409

    substituida = importar(cliente_autenticado, solicitacao_id, nova, substituir=True)
    assert substituida.status_code == 201
    assert substituida.json()["id"] == primeira.json()["id"]
    assert substituida.json()["valor_total"] == 9 * 10 + 2 * 4

    assert len(cliente_autenticado.get(f"/solicitacoes-compra/{solicitacao_id}/cotacoes").json()) == 1


def test_importar_erros_de_conteudo(cliente_autenticado):
    solicitacao_id, cabo, _ = criar_solicitacao_com_dois_itens(cliente_autenticado)
    outra_id = cliente_autenticado.post(
        "/solicitacoes-compra",
        json={"itens": [{"produto_id": cabo, "quantidade": 1}]}
    ).json()["id"]
    criar_fornecedor(cliente_autenticado, "Fornecedor A")

    sem_fornecedor = importar(
        cliente_autenticado, solicitacao_id,
        preencher(baixar_modelo(cliente_autenticado, solicitacao_id), "Não Existe", {cabo: 1})
    )
    assert sem_fornecedor.status_code == 404
    assert "Não Existe" in sem_fornecedor.json()["detail"]

    sem_precos = importar(
        cliente_autenticado, solicitacao_id,
        preencher(baixar_modelo(cliente_autenticado, solicitacao_id), "Fornecedor A", {})
    )
    assert sem_precos.status_code == 400

    de_outra = importar(
        cliente_autenticado, solicitacao_id,
        preencher(baixar_modelo(cliente_autenticado, outra_id), "Fornecedor A", {})
    )
    assert de_outra.status_code == 400
    assert f"Nº {outra_id}" in de_outra.json()["detail"]

    lixo = importar(cliente_autenticado, solicitacao_id, b"isto nao e uma planilha")
    assert lixo.status_code == 400


def test_consulta_nao_importa(cliente_autenticado):
    solicitacao_id, cabo, _ = criar_solicitacao_com_dois_itens(cliente_autenticado)

    resposta = importar(novo_cliente("CONSULTA"), solicitacao_id, b"x")

    assert resposta.status_code == 403


def test_exportar_e_baixar_planilha_da_cotacao(cliente_autenticado):
    solicitacao_id, cabo, conector = criar_solicitacao_com_dois_itens(cliente_autenticado)
    fornecedor_id = criar_fornecedor(cliente_autenticado, "Fornecedor A")
    cotacao = registrar_cotacao(cliente_autenticado, solicitacao_id, fornecedor_id, {cabo: 5, conector: 2}, frete=10)

    for formato, tipo in (("xlsx", XLSX), ("pdf", "application/pdf")):
        resposta = cliente_autenticado.get(
            f"/solicitacoes-compra/{solicitacao_id}/cotacoes/exportar",
            params={"formato": formato},
        )
        assert resposta.status_code == 200
        assert resposta.headers["content-type"] == tipo

    livro = load_workbook(BytesIO(
        cliente_autenticado.get(f"/solicitacoes-compra/{solicitacao_id}/cotacoes/exportar").content
    ))
    assert livro["Cotações"]["A2"].value == "Fornecedor A"

    planilha = cliente_autenticado.get(
        f"/solicitacoes-compra/{solicitacao_id}/cotacoes/{cotacao['id']}/planilha"
    )
    assert planilha.status_code == 200
    aba = load_workbook(BytesIO(planilha.content))["Orçamento"]
    assert campo(aba, "Fornecedor").value == "Fornecedor A"
    assert aba.cell(row=linhas_itens(aba)[cabo], column=4).value == 5
