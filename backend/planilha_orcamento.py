"""Planilha padrão de orçamento (cotação de fornecedor).

O modelo é gerado para uma solicitação de compra, já com os produtos e as
quantidades. O fornecedor (ou o comprador) preenche os campos amarelos e a
planilha volta para o sistema pela importação, virando uma cotação.

A leitura procura os rótulos da coluna A e o cabeçalho "Código" da tabela,
então linhas extras, cores ou larguras alteradas não atrapalham.
"""

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ABA = "Orçamento"
ABA_LISTAS = "Listas"

ROTULOS = {
    "solicitacao": "Solicitação Nº",
    "fornecedor": "Fornecedor",
    "forma_pagamento": "Forma de pagamento",
    "frete": "Frete (R$)",
    "prazo": "Prazo de entrega (dias)",
    "validade": "Validade da proposta",
    "observacao": "Observação",
}

CABECALHO_ITENS = ["Código", "Produto", "Quantidade", "Preço unitário (R$)", "Subtotal (R$)"]

FUNDO_TITULO = PatternFill("solid", fgColor="1C1745")
FUNDO_CABECALHO = PatternFill("solid", fgColor="263A44")
FUNDO_EDITAVEL = PatternFill("solid", fgColor="FFF4CC")
FUNDO_ROTULO = PatternFill("solid", fgColor="F4F5F8")
BORDA = Border(*(Side(style="thin", color="CDD1DC"),) * 4)
MOEDA = 'R$ #,##0.00'


class PlanilhaInvalida(ValueError):
    """Erro de conteúdo da planilha, com mensagem pronta para o usuário."""


def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().lower()


def _digitos(texto) -> str:
    return re.sub(r"\D", "", str(texto or ""))


def descrever_fornecedor(fornecedor) -> str:
    return fornecedor["nome"]


def descrever_forma(forma) -> str:
    return f"{forma['codigo']} - {forma['titulo']}"


def gerar_modelo(solicitacao, fornecedores, formas, cotacao=None) -> bytes:
    """Gera o .xlsx. Com `cotacao`, sai preenchido com os valores dela."""
    livro = Workbook()
    aba = livro.active
    aba.title = ABA

    listas = livro.create_sheet(ABA_LISTAS)
    nomes_fornecedores = [descrever_fornecedor(f) for f in fornecedores]
    nomes_formas = [descrever_forma(f) for f in formas]
    for linha, nome in enumerate(nomes_fornecedores, start=1):
        listas.cell(row=linha, column=1, value=nome)
    for linha, nome in enumerate(nomes_formas, start=1):
        listas.cell(row=linha, column=2, value=nome)
    listas.sheet_state = "hidden"

    aba.column_dimensions["A"].width = 26
    aba.column_dimensions["B"].width = 46
    aba.column_dimensions["C"].width = 14
    aba.column_dimensions["D"].width = 20
    aba.column_dimensions["E"].width = 18

    aba.merge_cells("A1:E1")
    titulo = aba["A1"]
    titulo.value = f"Orçamento — Solicitação de compra Nº {solicitacao['id']}"
    titulo.font = Font(bold=True, size=14, color="FFFFFF")
    titulo.fill = FUNDO_TITULO
    titulo.alignment = Alignment(vertical="center", indent=1)
    aba.row_dimensions[1].height = 30

    aba.merge_cells("A2:E2")
    aba["A2"].value = "Preencha os campos em amarelo e importe a planilha na solicitação. Deixe o preço em branco para itens sem oferta."
    aba["A2"].font = Font(italic=True, color="5F6577")

    precos = {}
    if cotacao:
        precos = {item["produto_id"]: item["preco_unitario"] for item in cotacao["itens"]}

    forma_atual = None
    if cotacao and cotacao.get("forma_pagamento_id"):
        forma_atual = next(
            (descrever_forma(f) for f in formas if f["id"] == cotacao["forma_pagamento_id"]),
            cotacao.get("forma_pagamento"),
        )

    valores = {
        "solicitacao": solicitacao["id"],
        "fornecedor": cotacao["fornecedor"] if cotacao else None,
        "forma_pagamento": forma_atual,
        "frete": cotacao["frete"] if cotacao else 0,
        "prazo": cotacao["prazo_entrega_dias"] if cotacao else None,
        "validade": cotacao["validade"] if cotacao else None,
        "observacao": cotacao["observacao"] if cotacao else None,
    }

    linha = 4
    linhas_campos = {}
    for chave, rotulo in ROTULOS.items():
        rotulo_celula = aba.cell(row=linha, column=1, value=rotulo)
        rotulo_celula.font = Font(bold=True)
        rotulo_celula.fill = FUNDO_ROTULO
        rotulo_celula.border = BORDA

        valor = aba.cell(row=linha, column=2, value=valores[chave])
        valor.border = BORDA
        if chave != "solicitacao":
            valor.fill = FUNDO_EDITAVEL
        if chave == "frete":
            valor.number_format = MOEDA
        if chave == "validade":
            valor.number_format = "DD/MM/YYYY"
        linhas_campos[chave] = linha
        linha += 1

    if nomes_fornecedores:
        validacao = DataValidation(
            type="list",
            formula1=f"={ABA_LISTAS}!$A$1:$A${len(nomes_fornecedores)}",
            allow_blank=True,
            showErrorMessage=False,
        )
        aba.add_data_validation(validacao)
        validacao.add(f"B{linhas_campos['fornecedor']}")

    if nomes_formas:
        validacao = DataValidation(
            type="list",
            formula1=f"={ABA_LISTAS}!$B$1:$B${len(nomes_formas)}",
            allow_blank=True,
        )
        aba.add_data_validation(validacao)
        validacao.add(f"B{linhas_campos['forma_pagamento']}")

    linha += 1
    linha_cabecalho = linha
    for coluna, texto in enumerate(CABECALHO_ITENS, start=1):
        celula = aba.cell(row=linha, column=coluna, value=texto)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = FUNDO_CABECALHO
        celula.border = BORDA

    for item in solicitacao["itens"]:
        linha += 1
        aba.cell(row=linha, column=1, value=item["produto_id"])
        aba.cell(row=linha, column=2, value=item["produto"])
        aba.cell(row=linha, column=3, value=item["quantidade"]).number_format = "#,##0"
        preco = aba.cell(row=linha, column=4, value=precos.get(item["produto_id"]))
        preco.number_format = MOEDA
        preco.fill = FUNDO_EDITAVEL
        subtotal = aba.cell(row=linha, column=5, value=f'=IF(D{linha}="","",C{linha}*D{linha})')
        subtotal.number_format = MOEDA
        for coluna in range(1, 6):
            aba.cell(row=linha, column=coluna).border = BORDA

    primeira, ultima = linha_cabecalho + 1, linha
    linha += 1
    aba.cell(row=linha, column=4, value="Itens").font = Font(bold=True)
    aba.cell(row=linha, column=5, value=f"=SUM(E{primeira}:E{ultima})").number_format = MOEDA
    linha += 1
    aba.cell(row=linha, column=4, value="Total com frete").font = Font(bold=True)
    total = aba.cell(row=linha, column=5, value=f"=E{linha - 1}+B{linhas_campos['frete']}")
    total.number_format = MOEDA
    total.font = Font(bold=True)

    aba.freeze_panes = f"A{linha_cabecalho + 1}"
    aba.sheet_view.showGridLines = False
    aba.print_title_rows = f"{linha_cabecalho}:{linha_cabecalho}"
    aba.page_setup.orientation = "portrait"
    aba.page_setup.fitToWidth = 1

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _numero(valor, campo: str, inteiro=False):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, (int, float)):
        numero = valor
    else:
        texto = str(valor).replace("R$", "").replace(" ", "").strip()
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        try:
            numero = float(texto)
        except ValueError:
            raise PlanilhaInvalida(f"Valor inválido em \"{campo}\": {valor}.")
    if numero < 0:
        raise PlanilhaInvalida(f"\"{campo}\" não pode ser negativo.")
    if inteiro:
        if float(numero) != int(numero):
            raise PlanilhaInvalida(f"\"{campo}\" deve ser um número inteiro.")
        return int(numero)
    return round(float(numero), 2)


def _data(valor):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise PlanilhaInvalida(f"Data inválida em \"{ROTULOS['validade']}\": {valor}. Use dd/mm/aaaa.")


def ler_planilha(conteudo: bytes):
    """Lê o .xlsx e devolve os campos crus (sem consultar o banco)."""
    try:
        livro = load_workbook(BytesIO(conteudo), data_only=True)
    except Exception:
        raise PlanilhaInvalida("Não foi possível ler o arquivo. Envie a planilha .xlsx do modelo de orçamento.")

    aba = livro[ABA] if ABA in livro.sheetnames else livro.worksheets[0]

    rotulos = {_normalizar(texto): chave for chave, texto in ROTULOS.items()}
    campos = {}
    linha_cabecalho = None

    for linha in aba.iter_rows(min_row=1, max_row=aba.max_row):
        primeira = linha[0].value
        chave = rotulos.get(_normalizar(primeira))
        if chave and len(linha) > 1:
            campos[chave] = linha[1].value
        if _normalizar(primeira) == _normalizar(CABECALHO_ITENS[0]):
            linha_cabecalho = linha[0].row
            break

    if linha_cabecalho is None:
        raise PlanilhaInvalida("A planilha não tem a tabela de itens (cabeçalho \"Código\"). Use o modelo do sistema.")

    itens = []
    for linha in aba.iter_rows(min_row=linha_cabecalho + 1, max_col=4, values_only=True):
        codigo, produto, _, preco = (list(linha) + [None] * 4)[:4]
        if codigo is None or str(codigo).strip() == "":
            continue
        try:
            produto_id = int(float(codigo))
        except (TypeError, ValueError):
            continue
        valor = _numero(preco, f"Preço unitário de {produto or produto_id}")
        if valor is not None:
            itens.append({"produto_id": produto_id, "preco_unitario": valor})

    return {
        "solicitacao": campos.get("solicitacao"),
        "fornecedor": campos.get("fornecedor"),
        "forma_pagamento": campos.get("forma_pagamento"),
        "frete": _numero(campos.get("frete"), ROTULOS["frete"]) or 0,
        "prazo": _numero(campos.get("prazo"), ROTULOS["prazo"], inteiro=True),
        "validade": _data(campos.get("validade")),
        "observacao": (str(campos["observacao"]).strip()[:500] or None) if campos.get("observacao") else None,
        "itens": itens,
    }


def encontrar_fornecedor(texto, fornecedores):
    alvo = _normalizar(texto)
    digitos = _digitos(texto)
    for fornecedor in fornecedores:
        if _normalizar(fornecedor["nome"]) == alvo:
            return fornecedor
        if digitos and len(digitos) >= 11 and _digitos(fornecedor.get("cpf_cnpj")) == digitos:
            return fornecedor
    return None


def encontrar_forma(texto, formas):
    alvo = _normalizar(texto)
    for forma in formas:
        if alvo in (_normalizar(descrever_forma(forma)), _normalizar(forma["codigo"]), _normalizar(forma["titulo"])):
            return forma
    return None
