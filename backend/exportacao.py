"""Geração de arquivos Excel (.xlsx) e PDF a partir de tabelas simples.

Cada tabela é descrita por colunas no formato (título, tipo, largura), onde
tipo é "texto", "inteiro", "moeda" ou "data", e por linhas com os valores.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class Tabela:
    titulo: str
    colunas: list[tuple[str, str, int]]
    linhas: list[list]
    observacao: str | None = None
    totais: list | None = field(default=None)


FORMATOS_EXCEL = {
    "moeda": 'R$ #,##0.00',
    "inteiro": '#,##0',
    "data": "DD/MM/YYYY",
}


def _nome_aba(titulo: str) -> str:
    proibidos = '[]:*?/\\'
    limpo = "".join("-" if c in proibidos else c for c in titulo)
    return limpo[:31] or "Dados"


def gerar_xlsx(tabelas: list[Tabela]) -> bytes:
    livro = Workbook()
    livro.remove(livro.active)

    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="263A44")

    for tabela in tabelas:
        aba = livro.create_sheet(_nome_aba(tabela.titulo))

        aba.append([titulo for titulo, _, _ in tabela.colunas])

        for celula in aba[1]:
            celula.font = cabecalho_fonte
            celula.fill = cabecalho_fundo
            celula.alignment = Alignment(vertical="center")

        for linha in tabela.linhas:
            aba.append(list(linha))

        if tabela.totais:
            aba.append(list(tabela.totais))
            for celula in aba[aba.max_row]:
                celula.font = Font(bold=True)

        for indice, (_, tipo, largura) in enumerate(tabela.colunas, start=1):
            letra = get_column_letter(indice)
            aba.column_dimensions[letra].width = max(10, largura)

            formato = FORMATOS_EXCEL.get(tipo)

            if formato:
                for (celula,) in aba.iter_rows(
                    min_row=2,
                    min_col=indice,
                    max_col=indice
                ):
                    celula.number_format = formato

        aba.freeze_panes = "A2"

        if tabela.linhas:
            aba.auto_filter.ref = (
                f"A1:{get_column_letter(len(tabela.colunas))}"
                f"{len(tabela.linhas) + 1}"
            )

    saida = BytesIO()
    livro.save(saida)

    return saida.getvalue()


def _texto_pdf(valor, tipo: str) -> str:
    if valor is None or valor == "":
        return "-"

    if tipo == "moeda":
        texto = f"{float(valor):,.2f}"
        texto = texto.replace(",", "X").replace(".", ",").replace("X", ".")
        return f"R$ {texto}"

    if tipo == "inteiro":
        return f"{int(valor):,}".replace(",", ".")

    if tipo == "data" and isinstance(valor, (date, datetime)):
        return valor.strftime("%d/%m/%Y %H:%M" if isinstance(valor, datetime) else "%d/%m/%Y")

    return str(valor)


def _latin1(texto: str) -> str:
    """As fontes padrão do PDF só aceitam Latin-1; troca o resto."""
    substituicoes = {"—": "-", "–": "-", "“": '"', "”": '"', "’": "'", "…": "..."}

    for original, novo in substituicoes.items():
        texto = texto.replace(original, novo)

    return texto.encode("latin-1", "replace").decode("latin-1")


class _Documento(FPDF):
    def __init__(self, titulo: str, subtitulo: str | None):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.titulo_documento = _latin1(titulo)
        self.subtitulo_documento = _latin1(subtitulo or "")
        self.set_auto_page_break(auto=True, margin=14)
        self.set_margins(12, 12, 12)

    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 8, self.titulo_documento, new_x="LMARGIN", new_y="NEXT")

        self.set_font("Helvetica", "", 9)
        self.set_text_color(90, 100, 105)
        self.cell(0, 5, self.subtitulo_documento, new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(0, 0, 0)
        self.ln(3)

    def footer(self):
        self.set_y(-10)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(90, 100, 105)
        self.cell(0, 5, f"Página {self.page_no()} de {{nb}}", align="R")
        self.set_text_color(0, 0, 0)


def _ajustar(pdf: FPDF, texto: str, largura: float) -> str:
    if pdf.get_string_width(texto) <= largura - 2:
        return texto

    while texto and pdf.get_string_width(texto + "...") > largura - 2:
        texto = texto[:-1]

    return texto + "..."


def gerar_pdf(titulo: str, subtitulo: str | None, tabelas: list[Tabela]) -> bytes:
    pdf = _Documento(titulo, subtitulo)
    pdf.add_page()

    largura_util = pdf.w - pdf.l_margin - pdf.r_margin

    for indice_tabela, tabela in enumerate(tabelas):
        if indice_tabela > 0:
            pdf.ln(4)

        if len(tabelas) > 1:
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 7, _latin1(tabela.titulo), new_x="LMARGIN", new_y="NEXT")

        if tabela.observacao:
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(0, 5, _latin1(tabela.observacao), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)

        soma = sum(largura for _, _, largura in tabela.colunas) or 1
        larguras = [largura_util * largura / soma for _, _, largura in tabela.colunas]
        alinhamentos = [
            "R" if tipo in ("moeda", "inteiro") else "L"
            for _, tipo, _ in tabela.colunas
        ]

        def cabecalho():
            pdf.set_font("Helvetica", "B", 9)
            pdf.set_fill_color(38, 58, 68)
            pdf.set_text_color(255, 255, 255)

            for (titulo_coluna, _, _), largura, alinhamento in zip(tabela.colunas, larguras, alinhamentos):
                pdf.cell(largura, 7, _ajustar(pdf, _latin1(titulo_coluna), largura), border=0, align=alinhamento, fill=True)

            pdf.ln()
            pdf.set_text_color(0, 0, 0)
            pdf.set_font("Helvetica", "", 9)

        cabecalho()

        if not tabela.linhas:
            pdf.cell(0, 7, "Nenhum registro.", new_x="LMARGIN", new_y="NEXT")
            continue

        linhas = list(tabela.linhas)

        if tabela.totais:
            linhas.append(tabela.totais)

        for numero, linha in enumerate(linhas):
            if pdf.will_page_break(6):
                pdf.add_page()
                cabecalho()

            eh_total = tabela.totais is not None and numero == len(linhas) - 1

            pdf.set_font("Helvetica", "B" if eh_total else "", 9)
            pdf.set_fill_color(240, 244, 246)

            for valor, (_, tipo, _), largura, alinhamento in zip(linha, tabela.colunas, larguras, alinhamentos):
                texto = _ajustar(pdf, _latin1(_texto_pdf(valor, tipo)), largura)
                pdf.cell(largura, 6, texto, border=0, align=alinhamento, fill=numero % 2 == 1)

            pdf.ln()

    return bytes(pdf.output())
